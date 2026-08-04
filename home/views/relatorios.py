from datetime import date
from calendar import monthrange

from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q, F
from django.shortcuts import render, get_object_or_404

from home.models import Presenca, Turma, Docente, Aluno
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, Alignment

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm
from io import BytesIO
import matplotlib.pyplot as plt
from reportlab.lib.utils import ImageReader
from reportlab.lib import colors
import openpyxl

from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from django.core.paginator import Paginator


@login_required
def presenca_aluno_mensal(request):

    user = request.user
    escola = user.escola

    # =====================================
    # CONTROLE DE ACESSO
    # =====================================
    professor = None

    if user.role == "professor":
        professor = Docente.objects.filter(
            user=user,
            escola=escola,
        ).first()

        if not professor:
            return render(request, "errors/403.html", status=403)

    elif user.role not in ("diretor", "coordenador"):
        return render(request, "errors/403.html", status=403)

    hoje = date.today()

    # =====================================
    # FILTROS
    # =====================================
    ano = int(request.GET.get("ano", hoje.year))
    mes = request.GET.get("mes")
    turma_id = request.GET.get("turma")

    # =====================================
    # PERÍODO
    # =====================================
    if mes:
        mes = int(mes)
        _, ultimo_dia = monthrange(ano, mes)
        data_inicio = date(ano, mes, 1)
        data_fim = date(ano, mes, ultimo_dia)
        tipo_relatorio = "mensal"
        mes_atual = mes
    else:
        data_inicio = date(ano, 1, 1)
        data_fim = date(ano, 12, 31)
        tipo_relatorio = "anual"
        mes_atual = None

    # =====================================
    # BASE DE PRESENÇAS
    # =====================================
    presencas = Presenca.objects.filter(
        chamada__data__range=(data_inicio, data_fim),
        aluno__escola=escola,
    )

    # Apenas professores enxergam suas próprias chamadas
    if user.role == "professor":
        presencas = presencas.filter(chamada__professor=professor)

    if turma_id and turma_id != "None":
        presencas = presencas.filter(chamada__turma_id=turma_id)

    # =====================================
    # AGRUPAMENTO
    # =====================================
    resumo = list(
        presencas.values(
            "aluno_id",
            "aluno__nome",
            "aluno__turma_principal__nome",
        )
        .annotate(
            total_aulas=Count("id"),
            total_presentes=Count(
                "id",
                filter=Q(presente=True),
            ),
            total_ausentes=Count(
                "id",
                filter=Q(presente=False),
            ),
        )
        .annotate(percentual_presenca=F("total_presentes") * 100.0 / F("total_aulas"))
        .order_by("aluno__nome")
    )
    # =====================================
    # AJUSTES
    # =====================================
    for r in resumo:
        if r["total_aulas"] == 0:
            r["percentual_presenca"] = 0
        else:
            r["percentual_presenca"] = round(
                r["percentual_presenca"],
                1,
            )

    # =====================================
    # DASHBOARD
    # =====================================
    total_alunos = len(resumo)

    total_aulas = sum(r["total_aulas"] for r in resumo)

    media_presenca = (
        round(
            sum(r["percentual_presenca"] for r in resumo) / total_alunos,
            1,
        )
        if total_alunos
        else 0
    )

    alunos_risco = sum(1 for r in resumo if r["percentual_presenca"] < 75)

    # ==========================================
    # PAGINAÇÃO
    # ==========================================

    paginator = Paginator(
        resumo,
        25,
    )

    page_number = request.GET.get("page")

    resumo = paginator.get_page(page_number)

    pagina_atual = resumo.number

    inicio = max(
        pagina_atual - 2,
        1,
    )

    fim = min(
        pagina_atual + 2,
        paginator.num_pages,
    )

    page_range = range(
        inicio,
        fim + 1,
    )

    # =====================================
    # DADOS AUXILIARES
    # =====================================
    turmas = Turma.objects.filter(
        escola=escola,
    ).order_by("nome")

    meses = [
        {"valor": 1, "nome": "Janeiro"},
        {"valor": 2, "nome": "Fevereiro"},
        {"valor": 3, "nome": "Março"},
        {"valor": 4, "nome": "Abril"},
        {"valor": 5, "nome": "Maio"},
        {"valor": 6, "nome": "Junho"},
        {"valor": 7, "nome": "Julho"},
        {"valor": 8, "nome": "Agosto"},
        {"valor": 9, "nome": "Setembro"},
        {"valor": 10, "nome": "Outubro"},
        {"valor": 11, "nome": "Novembro"},
        {"valor": 12, "nome": "Dezembro"},
    ]

    return render(
        request,
        "pages/relatorios/presenca_aluno_mensal.html",
        {
            "resumo": resumo,
            "page_range": page_range,
            "turmas": turmas,
            "meses": meses,
            "mes_atual": mes_atual,
            "ano_atual": ano,
            "turma_selecionada": turma_id,
            "tipo_relatorio": tipo_relatorio,
            "total_alunos": total_alunos,
            "total_aulas": total_aulas,
            "media_presenca": media_presenca,
            "alunos_risco": alunos_risco,
        },
    )


@login_required
def export_presenca_aluno_mensal_excel(request):
    """
    Exporta para Excel o relatório de presença por aluno
    - Mensal (quando mês é informado)
    - Anual (quando mês NÃO é informado)
    """

    user = request.user
    escola = user.escola

    # ==========================================
    # CONTROLE DE ACESSO
    # ==========================================

    professor = Docente.objects.filter(user=user).first()

    if not professor and user.role not in ("diretor", "coordenador"):
        return render(request, "errors/403.html", status=403)

    hoje = date.today()

    # ==========================================
    # FILTROS
    # ==========================================

    ano = int(request.GET.get("ano", hoje.year))
    mes = request.GET.get("mes")
    turma_id = request.GET.get("turma")

    # ==========================================
    # PERÍODO
    # ==========================================

    if mes:
        mes = int(mes)
        _, ultimo_dia = monthrange(ano, mes)

        data_inicio = date(ano, mes, 1)
        data_fim = date(ano, mes, ultimo_dia)

        tipo_relatorio = "mensal"
        periodo_label = f"{mes:02d}/{ano}"

    else:
        data_inicio = date(ano, 1, 1)
        data_fim = date(ano, 12, 31)

        tipo_relatorio = "anual"
        periodo_label = f"Ano {ano}"

    # ==========================================
    # QUERY
    # ==========================================

    presencas = Presenca.objects.filter(
        chamada__data__range=(data_inicio, data_fim),
        aluno__escola=escola,
    )

    if user.role == "professor":
        presencas = presencas.filter(
            chamada__professor=professor,
        )

    if turma_id and turma_id != "None":
        presencas = presencas.filter(
            chamada__turma_id=turma_id,
        )

    resumo = (
        presencas.values(
            "aluno__nome",
            "aluno__turma_principal__nome",
        )
        .annotate(
            total_aulas=Count("id"),
            total_presentes=Count(
                "id",
                filter=Q(presente=True),
            ),
            total_ausentes=Count(
                "id",
                filter=Q(presente=False),
            ),
        )
        .annotate(
            percentual_presenca=(F("total_presentes") * 100.0 / F("total_aulas")),
        )
        .order_by(
            "aluno__turma_principal__nome",
            "aluno__nome",
        )
    )

    # ==========================================
    # RESUMO GERAL
    # ==========================================

    total_alunos = len(resumo)

    total_aulas = sum(r["total_aulas"] for r in resumo)

    total_presentes = sum(r["total_presentes"] for r in resumo)

    total_faltas = sum(r["total_ausentes"] for r in resumo)

    media_frequencia = (total_presentes * 100 / total_aulas) if total_aulas else 0

    # ==========================================
    # CRIA EXCEL
    # ==========================================

    wb = openpyxl.Workbook()
    ws = wb.active

    ws.title = "Presença Mensal" if tipo_relatorio == "mensal" else "Presença Anual"

    # ==========================================
    # ESTILOS
    # ==========================================

    azul = "2563EB"
    verde = "22C55E"
    vermelho = "EF4444"
    amarelo = "F59E0B"
    cinza = "F3F4F6"
    branco = "FFFFFF"

    titulo_font = Font(
        size=18,
        bold=True,
        color=branco,
    )

    subtitulo_font = Font(
        size=12,
        bold=True,
    )

    header_font = Font(
        bold=True,
        color=branco,
    )

    bold = Font(
        bold=True,
    )

    center = Alignment(
        horizontal="center",
        vertical="center",
    )

    thin = Side(
        border_style="thin",
        color="DDDDDD",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    fill_azul = PatternFill(
        "solid",
        fgColor=azul,
    )

    fill_cinza = PatternFill(
        "solid",
        fgColor=cinza,
    )

    fill_verde = PatternFill(
        "solid",
        fgColor=verde,
    )

    fill_amarelo = PatternFill(
        "solid",
        fgColor=amarelo,
    )

    fill_vermelho = PatternFill(
        "solid",
        fgColor=vermelho,
    )

    # ==========================================
    # CABEÇALHO
    # ==========================================

    ws.merge_cells("A1:G1")

    cell = ws["A1"]
    cell.value = escola.nome.upper()
    cell.fill = fill_azul
    cell.font = titulo_font
    cell.alignment = center

    ws.merge_cells("A2:G2")

    ws["A2"] = "RELATÓRIO GERAL DE FREQUÊNCIA"
    ws["A2"].font = Font(
        bold=True,
        size=14,
    )

    ws["A4"] = "Período"
    ws["B4"] = periodo_label

    ws["A5"] = "Emitido em"
    ws["B5"] = date.today().strftime("%d/%m/%Y")

    ws["A7"] = "Total de alunos"
    ws["B7"] = total_alunos

    ws["A8"] = "Total de aulas"
    ws["B8"] = total_aulas

    ws["A9"] = "Presenças"
    ws["B9"] = total_presentes

    ws["A10"] = "Faltas"
    ws["B10"] = total_faltas

    ws["A11"] = "Frequência média"
    ws["B11"] = round(media_frequencia, 1)

    for row in range(4, 12):
        ws[f"A{row}"].font = bold

    linha_tabela = 13

    wb = openpyxl.Workbook()
    ws = wb.active

    ws.title = "Presença Mensal" if tipo_relatorio == "mensal" else "Presença Anual"

    # ==========================================
    # ESTILOS
    # ==========================================

    azul = "2563EB"
    verde = "22C55E"
    vermelho = "EF4444"
    amarelo = "F59E0B"
    cinza = "F3F4F6"
    branco = "FFFFFF"

    titulo_font = Font(size=18, bold=True, color=branco)
    subtitulo_font = Font(size=13, bold=True)
    header_font = Font(bold=True, color=branco)
    bold = Font(bold=True)

    center = Alignment(horizontal="center", vertical="center")

    thin = Side(border_style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    fill_azul = PatternFill("solid", fgColor=azul)
    fill_cinza = PatternFill("solid", fgColor=cinza)
    fill_verde = PatternFill("solid", fgColor=verde)
    fill_amarelo = PatternFill("solid", fgColor=amarelo)
    fill_vermelho = PatternFill("solid", fgColor=vermelho)

    # ==========================================
    # CABEÇALHO
    # ==========================================

    ws.merge_cells("A1:G1")
    ws["A1"] = escola.nome.upper()
    ws["A1"].font = titulo_font
    ws["A1"].fill = fill_azul
    ws["A1"].alignment = center

    ws.merge_cells("A2:G2")
    ws["A2"] = "RELATÓRIO GERAL DE FREQUÊNCIA"
    ws["A2"].font = subtitulo_font

    ws["A4"] = "Período"
    ws["B4"] = periodo_label

    ws["A5"] = "Emitido em"
    ws["B5"] = date.today().strftime("%d/%m/%Y")

    ws["A7"] = "Total de alunos"
    ws["B7"] = total_alunos

    ws["A8"] = "Total de aulas"
    ws["B8"] = total_aulas

    ws["A9"] = "Presenças"
    ws["B9"] = total_presentes

    ws["A10"] = "Faltas"
    ws["B10"] = total_faltas

    ws["A11"] = "Frequência média"
    ws["B11"] = round(media_frequencia, 1)

    for row in range(4, 12):
        ws[f"A{row}"].font = bold

    # ==========================================
    # TABELA
    # ==========================================

    linha_inicio = 13

    headers = [
        "Aluno",
        "Turma",
        "Total de Aulas",
        "Presenças",
        "Faltas",
        "Frequência (%)",
        "Situação",
    ]

    for col, titulo in enumerate(headers, start=1):

        cell = ws.cell(
            row=linha_inicio,
            column=col,
            value=titulo,
        )

        cell.font = header_font
        cell.fill = fill_azul
        cell.alignment = center
        cell.border = border

    linha = linha_inicio + 1

    for r in resumo:

        percentual = round(r["percentual_presenca"], 1)

        if percentual >= 75:
            situacao = "Frequência Regular"
            cor = fill_verde

        elif percentual >= 60:
            situacao = "Atenção"
            cor = fill_amarelo

        else:
            situacao = "Risco de Reprovação"
            cor = fill_vermelho

        dados = [
            r["aluno__nome"],
            r["aluno__turma_principal__nome"] or "-",
            r["total_aulas"],
            r["total_presentes"],
            r["total_ausentes"],
            percentual,
            situacao,
        ]

        for col, valor in enumerate(dados, start=1):

            cell = ws.cell(
                row=linha,
                column=col,
                value=valor,
            )

            cell.border = border

            if col >= 3:
                cell.alignment = center

        ws.cell(row=linha, column=7).fill = cor

        if linha % 2 == 0:
            for col in range(1, 8):
                if col != 7:
                    ws.cell(row=linha, column=col).fill = fill_cinza

        linha += 1

    # ==========================================
    # FILTROS
    # ==========================================

    ws.freeze_panes = "A14"
    ws.auto_filter.ref = f"A13:G{linha-1}"

    # ==========================================
    # LARGURA DAS COLUNAS
    # ==========================================

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 25

    # ==========================================
    # RESPONSE
    # ==========================================

    if tipo_relatorio == "mensal":
        filename = f"presenca_alunos_{mes:02d}_{ano}.xlsx"
    else:
        filename = f"presenca_alunos_anual_{ano}.xlsx"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    wb.save(response)

    return response


@login_required
def pdf_presenca_aluno_mensal(request):
    user = request.user
    escola = user.escola
    professor = Docente.objects.filter(user=user).first()

    if not professor and user.role not in ("diretor", "coordenador"):
        return render(request, "errors/403.html", status=403)

    hoje = date.today()

    # ==========================================
    # FILTROS
    # ==========================================

    ano = int(request.GET.get("ano", hoje.year))
    mes = request.GET.get("mes")
    turma_id = request.GET.get("turma")

    # ==========================================
    # PERÍODO
    # ==========================================

    if mes:
        mes = int(mes)
        _, ultimo_dia = monthrange(ano, mes)

        data_inicio = date(ano, mes, 1)
        data_fim = date(ano, mes, ultimo_dia)

        titulo = "RELATÓRIO GERAL DE FREQUÊNCIA"
        periodo_label = f"{mes:02d}/{ano}"
        filename = f"presenca_alunos_{mes:02d}_{ano}.pdf"

    else:
        data_inicio = date(ano, 1, 1)
        data_fim = date(ano, 12, 31)

        titulo = "RELATÓRIO GERAL DE FREQUÊNCIA"
        periodo_label = f"Ano {ano}"
        filename = f"presenca_alunos_anual_{ano}.pdf"

    # ==========================================
    # PRESENÇAS
    # ==========================================

    presencas = Presenca.objects.filter(
        chamada__data__range=(data_inicio, data_fim),
        aluno__escola=escola,
    )

    if user.role == "professor":
        presencas = presencas.filter(
            chamada__professor=professor,
        )

    if turma_id and turma_id != "None":
        presencas = presencas.filter(
            chamada__turma_id=turma_id,
        )

    resumo = (
        presencas.values(
            "aluno__nome",
            "aluno__turma_principal__nome",
        )
        .annotate(
            total_aulas=Count("id"),
            presentes=Count("id", filter=Q(presente=True)),
            faltas=Count("id", filter=Q(presente=False)),
        )
        .annotate(
            percentual=F("presentes") * 100.0 / F("total_aulas"),
        )
        .order_by("aluno__nome")
    )

    # ==========================================
    # RESUMO GERAL
    # ==========================================

    total_alunos = len(resumo)

    total_aulas = sum(item["total_aulas"] for item in resumo)

    total_presentes = sum(item["presentes"] for item in resumo)

    total_faltas = sum(item["faltas"] for item in resumo)

    media_frequencia = (total_presentes * 100 / total_aulas) if total_aulas else 0
    # ==========================================
    # PDF
    # ==========================================

    response = HttpResponse(content_type="application/pdf")

    response["Content-Disposition"] = f'inline; filename="{filename}"'

    pdf = canvas.Canvas(
        response,
        pagesize=A4,
    )

    largura, altura = A4

    # ==========================================
    # CORES
    # ==========================================

    azul = colors.HexColor("#2563EB")
    cinza = colors.HexColor("#6B7280")
    cinza_claro = colors.HexColor("#E5E7EB")

    # ==========================================
    # CABEÇALHO
    # ==========================================

    pdf.setFillColor(azul)

    pdf.rect(
        0,
        altura - 2.2 * cm,
        largura,
        2.2 * cm,
        stroke=0,
        fill=1,
    )

    pdf.setFillColor(colors.white)

    pdf.setFont(
        "Helvetica-Bold",
        22,
    )

    pdf.drawString(
        2 * cm,
        altura - 1.4 * cm,
        escola.nome.upper(),
    )

    pdf.setFillColor(colors.black)

    pdf.setFont(
        "Helvetica-Bold",
        20,
    )

    pdf.drawString(
        2 * cm,
        altura - 3.3 * cm,
        titulo,
    )

    pdf.setStrokeColor(cinza_claro)

    pdf.line(
        2 * cm,
        altura - 3.6 * cm,
        largura - 2 * cm,
        altura - 3.6 * cm,
    )

    # ==========================================
    # DADOS DA ESCOLA
    # ==========================================

    pdf.setFillColor(cinza)

    pdf.setFont(
        "Helvetica",
        9,
    )

    endereco = (
        f"{escola.endereco}, {escola.numero} - "
        f"{escola.bairro} - "
        f"{escola.cidade}/{escola.estado}"
    )

    pdf.drawString(
        2 * cm,
        altura - 4.2 * cm,
        f"CNPJ: {escola.cnpj}",
    )

    pdf.drawString(
        8.5 * cm,
        altura - 4.2 * cm,
        f"Telefone: {escola.telefone}",
    )

    pdf.drawString(
        2 * cm,
        altura - 4.8 * cm,
        endereco,
    )

    pdf.drawString(
        2 * cm,
        altura - 5.4 * cm,
        escola.email or "",
    )

    # ==========================================
    # PERÍODO
    # ==========================================

    pdf.setFillColor(colors.black)

    pdf.setFont(
        "Helvetica-Bold",
        16,
    )

    pdf.drawString(
        2 * cm,
        altura - 6.6 * cm,
        f"Período: {periodo_label}",
    )

    # ==========================================
    # CARDS
    # ==========================================

    card_y = altura - 10.2 * cm

    card_w = 3.3 * cm
    card_h = 2.6 * cm
    espaco = 0.35 * cm

    cards = [
        ("Alunos", total_alunos, "#2563EB"),
        ("Aulas", total_aulas, "#22C55E"),
        ("Presenças", total_presentes, "#F59E0B"),
        ("Média", f"{round(media_frequencia,1)}%", "#8B5CF6"),
    ]

    for i, (titulo_card, valor, cor) in enumerate(cards):

        x = 2 * cm + (card_w + espaco) * i

        pdf.setFillColor(colors.white)

        pdf.roundRect(
            x,
            card_y,
            card_w,
            card_h,
            8,
            stroke=1,
            fill=1,
        )

        pdf.setStrokeColor(colors.HexColor("#DDDDDD"))

        pdf.roundRect(
            x,
            card_y,
            card_w,
            card_h,
            8,
            stroke=1,
            fill=0,
        )

        pdf.setFillColor(colors.HexColor(cor))

        pdf.setFont(
            "Helvetica-Bold",
            16,
        )

        pdf.drawCentredString(
            x + card_w / 2,
            card_y + 1.45 * cm,
            str(valor),
        )

        pdf.setFillColor(cinza)

        pdf.setFont(
            "Helvetica",
            9,
        )

        pdf.drawCentredString(
            x + card_w / 2,
            card_y + 0.55 * cm,
            titulo_card,
        )

    # ==========================================
    # TABELA
    # ==========================================

    y = card_y - 1.8 * cm

    pdf.setFillColor(colors.black)

    pdf.setFont(
        "Helvetica-Bold",
        10,
    )

    pdf.drawString(2 * cm, y, "Aluno")
    pdf.drawString(8.2 * cm, y, "Turma")
    pdf.drawRightString(12.2 * cm, y, "Aulas")
    pdf.drawRightString(14.3 * cm, y, "Pres.")
    pdf.drawRightString(16.4 * cm, y, "Falt.")
    pdf.drawRightString(18.5 * cm, y, "%")

    pdf.setStrokeColor(cinza_claro)

    pdf.line(
        2 * cm,
        y - 0.2 * cm,
        largura - 2 * cm,
        y - 0.2 * cm,
    )

    y -= 0.8 * cm

    pdf.setFont(
        "Helvetica",
        9,
    )

    # ==========================================
    # LINHAS
    # ==========================================

    for r in resumo:

        percentual = round(r["percentual"], 1)

        pdf.setFillColor(colors.black)

        pdf.drawString(
            2 * cm,
            y,
            r["aluno__nome"][:30],
        )

        pdf.drawString(
            8.2 * cm,
            y,
            (r["aluno__turma_principal__nome"] or "-")[:15],
        )

        pdf.drawRightString(
            12.2 * cm,
            y,
            str(r["total_aulas"]),
        )

        pdf.drawRightString(
            14.3 * cm,
            y,
            str(r["presentes"]),
        )

        pdf.drawRightString(
            16.4 * cm,
            y,
            str(r["faltas"]),
        )

        if percentual >= 75:
            pdf.setFillColorRGB(0.16, 0.62, 0.35)
        elif percentual >= 60:
            pdf.setFillColorRGB(1, 0.60, 0)
        else:
            pdf.setFillColorRGB(0.80, 0.20, 0.20)

        pdf.drawRightString(
            18.5 * cm,
            y,
            f"{percentual}%",
        )

        y -= 0.6 * cm

        if y < 2.8 * cm:

            pdf.showPage()

            y = altura - 2.5 * cm

            pdf.setFont(
                "Helvetica-Bold",
                10,
            )

            pdf.drawString(2 * cm, y, "Aluno")
            pdf.drawString(8.2 * cm, y, "Turma")
            pdf.drawRightString(12.2 * cm, y, "Aulas")
            pdf.drawRightString(14.3 * cm, y, "Pres.")
            pdf.drawRightString(16.4 * cm, y, "Falt.")
            pdf.drawRightString(18.5 * cm, y, "%")

            pdf.line(
                2 * cm,
                y - 0.2 * cm,
                largura - 2 * cm,
                y - 0.2 * cm,
            )

            pdf.setFont(
                "Helvetica",
                9,
            )

            y -= 0.8 * cm

    # ==========================================
    # RESUMO
    # ==========================================

    if y > 5 * cm:

        y -= 0.8 * cm

        pdf.setFillColor(colors.black)

        pdf.setFont(
            "Helvetica-Bold",
            11,
        )

        pdf.drawString(
            2 * cm,
            y,
            "Resumo",
        )

        pdf.setFont(
            "Helvetica",
            9,
        )

        pdf.drawString(
            2 * cm,
            y - 0.6 * cm,
            (
                f"No período analisado foram registradas "
                f"{total_aulas} aulas para {total_alunos} alunos, "
                f"totalizando {total_presentes} presenças e "
                f"{total_faltas} faltas, com frequência média "
                f"de {round(media_frequencia,1)}%."
            ),
        )

    # ==========================================
    # RODAPÉ
    # ==========================================

    pdf.setStrokeColor(cinza_claro)

    pdf.line(
        2 * cm,
        2.3 * cm,
        largura - 2 * cm,
        2.3 * cm,
    )

    pdf.setFillColor(cinza)

    pdf.setFont(
        "Helvetica",
        8,
    )

    pdf.drawString(
        2 * cm,
        1.7 * cm,
        f"Documento emitido automaticamente em {date.today().strftime('%d/%m/%Y')}",
    )

    pdf.drawRightString(
        largura - 2 * cm,
        1.7 * cm,
        escola.nome,
    )

    pdf.showPage()
    pdf.save()

    return response


@login_required
def pdf_presenca_aluno_individual(request, aluno_id):
    user = request.user
    escola = user.escola

    professor = Docente.objects.filter(user=user).first()

    if not professor and user.role not in ("diretor", "coordenador"):
        return render(request, "errors/403.html", status=403)

    hoje = date.today()

    mes = request.GET.get("mes")
    ano = int(request.GET.get("ano", hoje.year))
    turma_id = request.GET.get("turma")

    # ==========================================
    # PERÍODO
    # ==========================================

    if mes:
        mes = int(mes)
        _, ultimo_dia = monthrange(ano, mes)

        data_inicio = date(ano, mes, 1)
        data_fim = date(ano, mes, ultimo_dia)

        titulo_periodo = f"{mes:02d}/{ano}"

    else:
        data_inicio = date(ano, 1, 1)
        data_fim = date(ano, 12, 31)

        titulo_periodo = f"Ano {ano}"

    aluno = get_object_or_404(
        Aluno,
        id=aluno_id,
        escola=escola,
    )

    # ==========================================
    # PRESENÇAS
    # ==========================================

    presencas = Presenca.objects.filter(
        aluno=aluno,
        chamada__data__range=(data_inicio, data_fim),
    )

    if user.role == "professor":
        presencas = presencas.filter(chamada__professor=professor)

    if turma_id and turma_id != "None":
        presencas = presencas.filter(chamada__turma_id=turma_id)

    total_aulas = presencas.count()
    presentes = presencas.filter(presente=True).count()

    faltas = presencas.filter(presente=False).count()

    percentual = presentes * 100 / total_aulas if total_aulas else 0

    # ==========================================
    # STATUS
    # ==========================================

    if percentual >= 75:
        status = "Frequência Regular"
        status_color = (0.16, 0.62, 0.35)

    elif percentual >= 60:
        status = "Atenção"
        status_color = (1, 0.60, 0)

    else:
        status = "Risco de Reprovação"
        status_color = (0.80, 0.20, 0.20)

    # ==========================================
    # GRÁFICO
    # ==========================================

    fig, ax = plt.subplots(figsize=(5.3, 5.3))

    if total_aulas == 0:

        ax.pie(
            [1],
            colors=["#d9d9d9"],
            startangle=90,
            wedgeprops=dict(
                width=0.4,
                edgecolor="white",
            ),
        )

        ax.text(
            0,
            0,
            "Sem\nDados",
            ha="center",
            va="center",
            fontsize=15,
            fontweight="bold",
        )

    else:

        ax.pie(
            [presentes, faltas],
            colors=[
                "#2ecc71",
                "#e74c3c",
            ],
            startangle=90,
            autopct=None,
            wedgeprops=dict(
                width=0.32,
                edgecolor="white",
                linewidth=2,
            ),
        )

        ax.text(
            0,
            0,
            f"{round(percentual, 1)}%",
            ha="center",
            va="center",
            fontsize=22,
            fontweight="bold",
            color="#111827",
        )

    plt.tight_layout()

    img_buffer = BytesIO()

    plt.savefig(
        img_buffer,
        format="png",
        transparent=True,
    )

    plt.close(fig)

    img_buffer.seek(0)

    # ==========================================
    # PDF
    # ==========================================

    response = HttpResponse(content_type="application/pdf")

    response["Content-Disposition"] = (
        f'inline; filename="frequencia_{aluno.nome.replace(" ", "_")}.pdf"'
    )

    pdf = canvas.Canvas(
        response,
        pagesize=A4,
    )

    largura, altura = A4

    # ==========================================
    # CORES
    # ==========================================

    azul = colors.HexColor("#2563EB")
    cinza = colors.HexColor("#6B7280")
    cinza_claro = colors.HexColor("#E5E7EB")

    # ==========================================
    # CABEÇALHO
    # ==========================================

    pdf.setFillColor(azul)

    pdf.rect(
        0,
        altura - 2.2 * cm,
        largura,
        2.2 * cm,
        stroke=0,
        fill=1,
    )

    pdf.setFillColor(colors.white)

    pdf.setFont(
        "Helvetica-Bold",
        22,
    )

    pdf.drawString(
        2 * cm,
        altura - 1.4 * cm,
        escola.nome.upper(),
    )

    pdf.setFillColor(colors.black)

    pdf.setFont(
        "Helvetica-Bold",
        20,
    )

    pdf.drawString(
        2 * cm,
        altura - 3.3 * cm,
        "RELATÓRIO INDIVIDUAL DE FREQUÊNCIA",
    )

    pdf.setStrokeColor(cinza_claro)

    pdf.line(
        2 * cm,
        altura - 3.6 * cm,
        largura - 2 * cm,
        altura - 3.6 * cm,
    )

    # ==========================================
    # DADOS DA ESCOLA
    # ==========================================

    pdf.setFillColor(cinza)

    pdf.setFont(
        "Helvetica",
        9,
    )

    endereco = (
        f"{escola.endereco}, {escola.numero} - "
        f"{escola.bairro} - "
        f"{escola.cidade}/{escola.estado}"
    )

    pdf.drawString(
        2 * cm,
        altura - 4.2 * cm,
        f"CNPJ: {escola.cnpj}",
    )

    pdf.drawString(
        8.5 * cm,
        altura - 4.2 * cm,
        f"Telefone: {escola.telefone}",
    )

    pdf.drawString(
        2 * cm,
        altura - 4.8 * cm,
        endereco,
    )

    pdf.drawString(
        2 * cm,
        altura - 5.4 * cm,
        escola.email or "",
    )

    # ==========================================
    # DADOS DO ALUNO
    # ==========================================

    y = altura - 6.8 * cm

    pdf.setFillColor(colors.black)

    pdf.setFont(
        "Helvetica-Bold",
        16,
    )

    pdf.drawString(
        2 * cm,
        y,
        aluno.nome,
    )

    pdf.setFont(
        "Helvetica",
        10,
    )

    pdf.setFillColor(cinza)

    pdf.drawString(
        2 * cm,
        y - 0.6 * cm,
        f"Turma: {aluno.turma_principal.nome if aluno.turma_principal else '-'}",
    )

    pdf.drawString(
        9 * cm,
        y - 0.6 * cm,
        f"Período: {titulo_periodo}",
    )

    # ==========================================
    # CARDS
    # ==========================================

    card_y = altura - 10.5 * cm

    card_w = 3.3 * cm
    card_h = 2.6 * cm
    espaco = 0.35 * cm

    cards = [
        ("Aulas", total_aulas, "#2563EB"),
        ("Presentes", presentes, "#22C55E"),
        ("Faltas", faltas, "#EF4444"),
        ("Frequência", f"{round(percentual, 1)}%", "#F59E0B"),
    ]

    for i, (titulo, valor, cor) in enumerate(cards):
        x = 2 * cm + (card_w + espaco) * i

        pdf.setFillColor(colors.white)

        pdf.roundRect(
            x,
            card_y,
            card_w,
            card_h,
            8,
            stroke=1,
            fill=1,
        )

        pdf.setStrokeColor(colors.HexColor("#DDDDDD"))

        pdf.roundRect(
            x,
            card_y,
            card_w,
            card_h,
            8,
            stroke=1,
            fill=0,
        )

        pdf.setFillColor(colors.HexColor(cor))

        pdf.setFont(
            "Helvetica-Bold",
            16,
        )

        pdf.drawCentredString(
            x + card_w / 2,
            card_y + 1.45 * cm,
            str(valor),
        )

        pdf.setFillColor(cinza)

        pdf.setFont(
            "Helvetica",
            9,
        )

        pdf.drawCentredString(
            x + card_w / 2,
            card_y + 0.55 * cm,
            titulo,
        )

        # ==========================================
    # STATUS
    # ==========================================

    status_y = card_y - 2.9 * cm

    pdf.setFillColor(colors.HexColor("#F8FAFC"))

    pdf.roundRect(
        2 * cm,
        status_y,
        8.5 * cm,
        2.2 * cm,
        8,
        stroke=0,
        fill=1,
    )

    pdf.setFillColor(cinza)

    pdf.setFont(
        "Helvetica-Bold",
        10,
    )

    pdf.drawString(
        3.4 * cm,
        status_y + 1.5 * cm,
        "SITUAÇÃO DO ALUNO",
    )

    pdf.setFillColorRGB(*status_color)

    pdf.setFont(
        "Helvetica-Bold",
        14,
    )

    pdf.drawString(
        2.0 * cm,
        status_y + 0.5 * cm,
        status.upper(),
    )

    # ==========================================
    # GRÁFICO
    # ==========================================

    pdf.setFont(
        "Helvetica-Bold",
        12,
    )

    image = ImageReader(img_buffer)

    pdf.setFillColor(colors.black)

    pdf.drawImage(
        image,
        11.2 * cm,
        altura - 18.8 * cm,
        width=8.4 * cm,
        height=8.4 * cm,
        mask="auto",
    )

    resumo_y = 4.2 * cm

    pdf.setFillColor(colors.black)

    pdf.setFont(
        "Helvetica-Bold",
        11,
    )

    pdf.drawString(2 * cm, resumo_y + 1.2 * cm, "Resumo")

    pdf.setFont(
        "Helvetica",
        9,
    )

    texto = (
        f"No período analisado o aluno participou de "
        f"{total_aulas} aulas, registrando presença em "
        f"{presentes} delas, alcançando frequência de "
        f"{round(percentual,1)}%."
    )

    pdf.drawString(
        2 * cm,
        resumo_y + 0.5 * cm,
        texto,
    )

    # ==========================================
    # RODAPÉ
    # ==========================================

    pdf.setStrokeColor(cinza_claro)

    pdf.line(
        2 * cm,
        2.3 * cm,
        largura - 2 * cm,
        2.3 * cm,
    )

    pdf.setFillColor(cinza)

    pdf.setFont(
        "Helvetica",
        8,
    )

    pdf.drawString(
        2 * cm,
        1.7 * cm,
        f"Documento emitido automaticamente em {date.today().strftime('%d/%m/%Y')}",
    )

    pdf.drawRightString(
        largura - 2 * cm,
        1.7 * cm,
        escola.nome,
    )

    # Finaliza a página
    pdf.showPage()

    # Grava o PDF
    pdf.save()

    return response
